from typing import Iterable, Mapping, Tuple
from io import BytesIO

from openpyxl import Workbook

from docx import Document

from fpdf import FPDF 

def export_xlsx(rows: Iterable[Mapping], filename: str) -> Tuple[bytes, str, str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    rows = list(rows)
    headers = list(rows[0].keys()) if rows else []
    if headers:
        ws.append(headers)
        for r in rows:
            ws.append([r.get(h) for h in headers])
    bio = BytesIO()
    wb.save(bio); bio.seek(0)
    return bio.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{filename}.xlsx"

# a .docx
def export_docx(rows: Iterable[Mapping], title: str, filename: str) -> Tuple[bytes, str, str]:
    doc = Document()
    doc.add_heading(title, level=1)
    rows = list(rows)
    headers = list(rows[0].keys()) if rows else []
    if headers:
        table = doc.add_table(rows=1, cols=len(headers))
        hdr = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = str(h)
        for r in rows:
            row = table.add_row().cells
            for i, h in enumerate(headers):
                v = r.get(h)
                row[i].text = "" if v is None else str(v)
    bio = BytesIO()
    doc.save(bio); bio.seek(0)
    return bio.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{filename}.docx"    


# a pdf
def _to_latin1(text) -> str:
    """
    fpdf2 con fuentes core usa Latin-1; forzamos conversión.
    (á, é, í, ó, ú, ñ, Ñ están cubiertos por Latin-1).
    """
    if text is None:
        return ""
    s = str(text)
    try:
        s.encode("latin-1")
        return s
    except UnicodeEncodeError:
        return s.encode("latin-1", "replace").decode("latin-1")

def export_pdf(rows: Iterable[Mapping], title: str, filename: str) -> Tuple[bytes, str, str]:
    """
    Genera un PDF simple en A4 apaisado con título y tabla.
    Devuelve (bytes, content_type, suggested_filename).
    """
    rows = list(rows)
    headers = list(rows[0].keys()) if rows else []

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    # Título
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, _to_latin1(title), ln=1)

    # Si no hay datos, mensaje simple
    if not headers:
        pdf.set_font("Helvetica", "", 12)
        pdf.cell(0, 8, _to_latin1("(sin datos)"), ln=1)
        out = pdf.output(dest="S")
        return (out if isinstance(out, (bytes, bytearray)) else out.encode("latin-1"),
                "application/pdf", f"{filename}.pdf")

    # Cálculo de ancho por columna
    left_margin = pdf.l_margin
    right_margin = pdf.r_margin
    page_width = pdf.w - left_margin - right_margin
    ncols = max(1, len(headers))
    col_w = page_width / ncols

    # Encabezados
    pdf.set_font("Helvetica", "B", 10)
    for h in headers:
        pdf.cell(col_w, 8, _to_latin1(h), border=1, ln=0, align="L")
    pdf.ln(8)

    # Filas (valores)
    pdf.set_font("Helvetica", "", 9)
    row_h = 6
    for r in rows:
        for h in headers:
            val = _to_latin1(r.get(h, ""))
            pdf.cell(col_w, row_h, val, border=1, ln=0, align="L")
        pdf.ln(row_h)

    out = pdf.output(dest="S")
    return (out if isinstance(out, (bytes, bytearray)) else out.encode("latin-1"),
            "application/pdf", f"{filename}.pdf")